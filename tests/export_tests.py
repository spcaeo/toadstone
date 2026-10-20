"""Tests for the exports API (src/api/exports.py).

Drives the real Flask app with its test client against a throwaway database,
so no server and no network are needed:

    python tests/export_tests.py

Each check pins a promise the export makes to whoever receives the file: every
dataset actually carries rows, the formats open in the tools people use (Excel
included), the filters they were looking at come with them, and the shareable
report never hands somebody else's HTML to the reader's browser.
"""

import csv
import io
import json
import os
import sys
import tempfile
import zipfile

os.environ.setdefault('LOCAL_MODE', 'true')

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src import crawl_db  # noqa: E402
from src.api.exports import DATASETS, create_blueprint  # noqa: E402
from src.utils.issue_factory import make_issue  # noqa: E402

PAGE_COUNT = 5  # the fixture's crawled pages, including one image row

PASSED, FAILED = [], []

# A title only a crawled page would carry: it must never reach the report as
# markup, whichever format the report is asked for.
HOSTILE_TITLE = 'Pricing <script>alert("xss")</script>'


def result(name, ok, detail=''):
    (PASSED if ok else FAILED).append(name)
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{(' - ' + detail) if detail else ''}")


def build_fixture_crawls():
    """Two crawls of one site, seeded so every dataset has something to say."""
    first = crawl_db.create_crawl(user_id=1, session_id='s1', base_url='https://site.test',
                                  base_domain='site.test', config_snapshot={})
    second = crawl_db.create_crawl(user_id=1, session_id='s2', base_url='https://site.test',
                                   base_domain='site.test', config_snapshot={})

    pages = [
        {'url': 'https://site.test/', 'status_code': 200, 'title': 'Home', 'h1': 'Home',
         'content_type': 'text/html', 'is_internal': True, 'depth': 0,
         'meta_description': 'The home page', 'canonical_url': 'https://site.test/',
         'word_count': 400, 'main_text': 'shared body', 'main_text_hash': 'dup',
         'headings': [{'level': 1, 'text': 'Home', 'is_visible': True, 'dom_order': 0},
                      {'level': 2, 'text': 'What we do', 'is_visible': True, 'dom_order': 1}],
         'heading_text_issues': [{'dom_order': 1, 'level': 2, 'problem': 'double space'}],
         'images': [{'src': 'https://site.test/logo.png', 'alt': 'Logo', 'has_alt': True,
                     'width': '200', 'height': '60'},
                    {'src': 'https://site.test/hero.jpg', 'alt': '', 'has_alt': False,
                     'width': '', 'height': ''}]},
        {'url': 'https://site.test/pricing', 'status_code': 200, 'title': HOSTILE_TITLE,
         'h1': 'Pricing', 'content_type': 'text/html', 'is_internal': True, 'depth': 1,
         'meta_description': '', 'canonical_url': 'https://site.test/pricing',
         'word_count': 120, 'main_text': 'shared body', 'main_text_hash': 'dup',
         'headings': [{'level': 1, 'text': 'Pricing', 'is_visible': True, 'dom_order': 0}],
         'images': [{'src': 'https://site.test/table.png', 'alt': 'Plans', 'has_alt': True,
                     'width': '', 'height': ''}]},
        # A page reached through a redirect, so the redirects view has a chain
        {'url': 'https://site.test/new', 'status_code': 200, 'title': 'Moved here',
         'content_type': 'text/html', 'is_internal': True, 'depth': 1, 'word_count': 300,
         'main_text_hash': 'moved',
         'redirects': [{'url': 'https://site.test/old', 'status_code': 301}]},
        # A page that 404s, so the broken-links view has a target
        {'url': 'https://site.test/gone', 'status_code': 404, 'title': '',
         'content_type': 'text/html', 'is_internal': True, 'depth': 1},
        {'url': 'https://site.test/logo.png', 'status_code': 200, 'title': '',
         'content_type': 'image/png', 'is_internal': True, 'depth': 1},
    ]
    links = [
        {'source_url': 'https://site.test/', 'target_url': 'https://site.test/pricing',
         'anchor_text': 'See pricing', 'is_internal': True, 'target_status': 200,
         'target_domain': 'site.test', 'placement': 'body', 'rel': '', 'nofollow': False,
         'link_target': '', 'in_nav': False, 'in_main': True, 'in_heading': False},
        {'source_url': 'https://site.test/', 'target_url': 'https://site.test/gone',
         'anchor_text': 'Old offer', 'is_internal': True, 'target_status': 404,
         'target_domain': 'site.test', 'placement': 'body', 'rel': '', 'nofollow': False,
         'link_target': '', 'in_nav': False, 'in_main': True, 'in_heading': False},
        {'source_url': 'https://site.test/pricing', 'target_url': 'https://other.test/x',
         'anchor_text': 'partner', 'is_internal': False, 'target_status': 200,
         'target_domain': 'other.test', 'placement': 'body', 'rel': 'nofollow',
         'nofollow': True, 'link_target': '_blank', 'in_nav': False, 'in_main': True,
         'in_heading': False},
    ]

    thin = make_issue('https://site.test/pricing', 'thin_content', 'warning', 'Content',
                      'Thin Content', 'Page has only 120 words',
                      offending_value='120', evidence={'found': '120 words'},
                      how_to_fix='Add substantive copy.')
    missing_desc = make_issue('https://site.test/pricing', 'missing_meta_description', 'error',
                              'SEO', 'Missing Meta Description', 'No meta description',
                              evidence={'selector': 'head > meta[name=description]'},
                              how_to_fix='Write one.')
    slow = make_issue('https://site.test/', 'slow_response_time', 'warning', 'Performance',
                      'Slow Response Time', 'Took 4000ms', offending_value='4000')
    broken = make_issue('https://site.test/', 'links_to_broken_urls', 'error', 'Links',
                        'Broken Link', 'Links to a page that 404s',
                        evidence={'related_urls': ['https://site.test/gone']})

    for crawl_id, issues in ((first, [thin, missing_desc]),
                             (second, [thin, missing_desc, slow, broken])):
        crawl_db.save_url_batch(crawl_id, pages)
        crawl_db.save_links_batch(crawl_id, links)
        crawl_db.save_issues_batch(crawl_id, issues)
        crawl_db.set_crawl_status(crawl_id, 'completed')

    crawl_db.save_sitemap_audit(second, [{'url': 'https://site.test/sitemap.xml', 'status': 200,
                                          'entries': 2, 'bytes': 300, 'is_index': False,
                                          'parent': '', 'error': '', 'error_kind': ''}],
                                {'in_both': ['https://site.test/'],
                                 'crawled_only': ['https://site.test/pricing'],
                                 'sitemap_only': ['https://site.test/missing']})
    return first, second, {'thin': thin, 'missing_desc': missing_desc, 'slow': slow}


def client():
    import main
    main.app.config['TESTING'] = True
    return main.app.test_client()


def get_json(app, url):
    response = app.get(url)
    return response.status_code, json.loads(response.data or b'{}')


def export_url(dataset, fmt='json', extra=''):
    return f'/api/crawls/{SECOND}/export/{dataset}?format={fmt}{extra}'


def test_every_dataset_has_rows():
    app = client()
    for dataset in DATASETS:
        status, body = get_json(app, export_url(dataset))
        result(f'dataset {dataset}: returns rows',
               status == 200 and body.get('count', 0) > 0,
               f"status={status} count={body.get('count')}")


def test_dataset_columns_are_the_documented_ones():
    app = client()
    for dataset, spec in DATASETS.items():
        _, body = get_json(app, export_url(dataset))
        first = (body.get('rows') or [{}])[0]
        result(f'dataset {dataset}: columns match the catalogue',
               set(first) == set(spec['fields']),
               f"missing={sorted(set(spec['fields']) - set(first))}")


def test_catalogue_lists_every_dataset():
    app = client()
    status, body = get_json(app, '/api/export/datasets')
    listed = {entry['dataset'] for entry in body.get('datasets', [])}
    result('catalogue: lists every dataset with its columns',
           status == 200 and listed == set(DATASETS)
           and all(entry['columns'] for entry in body['datasets']),
           f'status={status} listed={len(listed)}')


def test_derived_views_are_built_from_the_stored_rows():
    app = client()
    _, dupes = get_json(app, export_url('duplicates'))
    urls = {row['url'] for row in dupes['rows']}
    result('duplicates: the two pages sharing a text hash are grouped',
           urls == {'https://site.test/', 'https://site.test/pricing'}
           and {row['group'] for row in dupes['rows']} == {1}, str(sorted(urls)))

    _, redirects = get_json(app, export_url('redirects'))
    result('redirects: the stored chain becomes one row per hop',
           redirects['count'] == 1
           and redirects['rows'][0]['from_url'] == 'https://site.test/old'
           and redirects['rows'][0]['status_code'] == 301
           and redirects['rows'][0]['to_url'] == 'https://site.test/new',
           str(redirects['rows']))

    _, broken = get_json(app, export_url('broken-links'))
    result('broken-links: only the link to the 404 is listed',
           broken['count'] == 1
           and broken['rows'][0]['target_url'] == 'https://site.test/gone'
           and broken['rows'][0]['status_code'] == 404, str(broken['rows']))

    _, headings = get_json(app, export_url('headings'))
    flagged = [row for row in headings['rows'] if row['problem']]
    result('headings: every heading is a row, carrying its text problem',
           headings['count'] == 3 and len(flagged) == 1
           and flagged[0]['level'] == 2, str(headings['rows']))

    _, images = get_json(app, export_url('images'))
    without_alt = [row for row in images['rows'] if not row['has_alt']]
    result('images: one row per img, keeping a missing alt apart from an empty one',
           images['count'] == 3 and len(without_alt) == 1
           and without_alt[0]['src'].endswith('hero.jpg'), str(images['rows']))

    _, coverage = get_json(app, export_url('sitemap-coverage'))
    states = {row['state'] for row in coverage['rows']}
    result('sitemap-coverage: each list becomes rows tagged with its state',
           coverage['count'] == 3
           and states == {'in_both', 'crawled_only', 'sitemap_only'}, str(sorted(states)))


def test_csv_format():
    app = client()
    response = app.get(export_url('issues', 'csv'))
    body = response.data.decode('utf-8-sig')
    rows = list(csv.DictReader(io.StringIO(body)))
    result('csv: parses with a header and one row per issue',
           response.status_code == 200 and len(rows) == 4 and 'evidence' in rows[0]
           and rows[0]['url'].startswith('https://'), f'rows={len(rows)}')
    result('csv: starts with the utf-8 BOM so Excel reads it as utf-8',
           response.data.startswith(b'\xef\xbb\xbf'), str(response.data[:6]))
    result('csv: is offered as a download with a dated filename',
           'attachment; filename=crawl-' in response.headers.get('Content-Disposition', '')
           and response.headers['Content-Disposition'].endswith('.csv'),
           response.headers.get('Content-Disposition', ''))
    result('csv: is served as text/csv',
           response.headers.get('Content-Type', '').startswith('text/csv'),
           response.headers.get('Content-Type', ''))


def test_ndjson_format():
    app = client()
    response = app.get(export_url('pages', 'ndjson'))
    lines = [line for line in response.data.decode().splitlines() if line.strip()]
    parsed = [json.loads(line) for line in lines]
    result('ndjson: one json object per line, one line per page',
           len(parsed) == PAGE_COUNT and all('url' in row for row in parsed),
           f'lines={len(lines)}')
    result('ndjson: served as application/x-ndjson',
           response.headers.get('Content-Type', '').startswith('application/x-ndjson'),
           response.headers.get('Content-Type', ''))


def test_field_selection():
    app = client()
    _, body = get_json(app, export_url('pages', 'json', '&fields=url,status_code'))
    result('fields: only the requested columns come back',
           all(set(row) == {'url', 'status_code'} for row in body['rows']),
           str(body['rows'][:1]))


def test_xlsx_format():
    from openpyxl import load_workbook
    app = client()
    response = app.get(export_url('issues', 'xlsx'))
    result('xlsx: served as a spreadsheet download',
           response.status_code == 200
           and 'spreadsheetml' in response.headers.get('Content-Type', '')
           and response.headers.get('Content-Disposition', '').endswith('.xlsx'),
           response.headers.get('Content-Type', ''))

    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.values)
    result('xlsx: opens with openpyxl, header row then one row per issue',
           workbook.sheetnames == ['issues'] and rows[0][0] == 'id' and len(rows) == 5,
           f'sheets={workbook.sheetnames} rows={len(rows)}')
    workbook.close()


def test_export_all_workbook():
    from openpyxl import load_workbook
    app = client()
    response = app.get(f'/api/crawls/{SECOND}/export/all?format=xlsx')
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    result('export/all xlsx: one sheet per dataset',
           response.status_code == 200 and workbook.sheetnames == list(DATASETS),
           str(workbook.sheetnames))
    populated = [name for name in workbook.sheetnames
                 if sum(1 for _ in workbook[name].values) > 1]
    result('export/all xlsx: every sheet carries rows under its header',
           populated == list(DATASETS),
           str(sorted(set(DATASETS) - set(populated))))
    workbook.close()


def test_export_all_zip():
    app = client()
    response = app.get(f'/api/crawls/{SECOND}/export/all?format=zip')
    result('export/all zip: served as a zip download',
           response.status_code == 200
           and response.headers.get('Content-Type', '').startswith('application/zip')
           and response.headers.get('Content-Disposition', '').endswith('.zip'),
           response.headers.get('Content-Type', ''))

    archive = zipfile.ZipFile(io.BytesIO(response.data))
    names = archive.namelist()
    result('export/all zip: one csv per dataset',
           len(names) == len(DATASETS)
           and all(name.endswith('.csv') for name in names)
           and all(any(f'-{dataset}-' in name for name in names) for dataset in DATASETS),
           str(names))

    body = archive.read(next(n for n in names if '-issues-' in n)).decode('utf-8-sig')
    rows = list(csv.DictReader(io.StringIO(body)))
    result('export/all zip: the csvs inside parse and carry their rows',
           len(rows) == 4 and rows[0]['url'].startswith('https://'), f'rows={len(rows)}')
    result('export/all zip: the archive is not corrupt', archive.testzip() is None)


def test_filters_narrow_the_export():
    app = client()
    _, all_issues = get_json(app, export_url('issues'))
    _, errors = get_json(app, export_url('issues', 'json', '&severity=error'))
    result('filter severity: narrows the issue export',
           errors['count'] == 2 and all_issues['count'] == 4
           and {row['type'] for row in errors['rows']} == {'error'},
           f"all={all_issues['count']} errors={errors['count']}")

    _, by_check = get_json(app, export_url('issues', 'json', '&check=thin_content'))
    result('filter check: narrows the issue export',
           [row['check'] for row in by_check['rows']] == ['thin_content'],
           str([row['check'] for row in by_check['rows']]))

    _, by_category = get_json(app, export_url('issues', 'json', '&category=Performance'))
    result('filter category: narrows the issue export',
           by_category['count'] == 1 and by_category['rows'][0]['category'] == 'Performance',
           str(by_category['count']))

    _, by_prefix = get_json(app, export_url('pages', 'json',
                                            '&url_prefix=https://site.test/pricing'))
    result('filter url_prefix: narrows the page export', by_prefix['count'] == 1,
           str(by_prefix['count']))

    _, by_status = get_json(app, export_url('pages', 'json', '&status_code=404'))
    result('filter status_code: narrows the page export',
           by_status['count'] == 1 and by_status['rows'][0]['url'].endswith('/gone'),
           str(by_status['count']))

    _, internal = get_json(app, export_url('links', 'json', '&internal=true'))
    _, external = get_json(app, export_url('links', 'json', '&internal=false'))
    result('filter internal: splits the link export',
           internal['count'] == 2 and external['count'] == 1,
           f"internal={internal['count']} external={external['count']}")

    csv_filtered = app.get(export_url('issues', 'csv', '&severity=error'))
    rows = list(csv.DictReader(io.StringIO(csv_filtered.data.decode('utf-8-sig'))))
    result('filter: the same narrowing applies to the csv download', len(rows) == 2,
           f'rows={len(rows)}')


def test_html_report():
    app = client()
    response = app.get(f'/api/crawls/{SECOND}/report?format=html')
    html = response.data.decode()
    result('report html: served as a self-contained page',
           response.status_code == 200
           and response.headers.get('Content-Type', '').startswith('text/html')
           and html.startswith('<!doctype html'), response.headers.get('Content-Type', ''))
    result('report html: pulls in no remote asset',
           'http://' not in html and 'src=' not in html and '<link' not in html)
    result('report html: carries the crawl metadata',
           'https://site.test' in html and f'Crawl {SECOND}' in html)
    result('report html: carries the counts by severity',
           '>2</div>' in html and 'Errors' in html and 'Warnings' in html)
    result('report html: names the top checks',
           'Thin Content' in html and 'Missing Meta Description' in html)
    result('report html: lists the worst pages',
           'https://site.test/pricing' in html and 'Worst pages' in html)
    result('report html: reports the sitemap coverage',
           'Sitemap coverage' in html and 'Listed but not crawled' in html)
    result('report html: reports what changed against the previous crawl',
           f'Compared against crawl {FIRST}' in html and 'Slow Response Time' in html)

    result('report html: escapes a page title that contains markup',
           '<script>' not in html and '&lt;script&gt;alert(&quot;xss&quot;)' in html)
    result('report html: is offered with a dated filename',
           'filename=crawl-' in response.headers.get('Content-Disposition', '')
           and response.headers['Content-Disposition'].endswith('.html'),
           response.headers.get('Content-Disposition', ''))


def test_markdown_report():
    app = client()
    response = app.get(f'/api/crawls/{SECOND}/report?format=md')
    text = response.data.decode()
    result('report md: served as markdown with a dated filename',
           response.status_code == 200
           and response.headers.get('Content-Type', '').startswith('text/markdown')
           and response.headers.get('Content-Disposition', '').endswith('.md'),
           response.headers.get('Content-Type', ''))
    result('report md: carries the metadata, the counts and the top checks',
           '# site.test' in text and '| Errors | 2 |' in text and 'Thin Content' in text)
    result('report md: reports what changed against the previous crawl',
           f'Compared against crawl {FIRST}' in text and 'Slow Response Time' in text)
    result('report md: a hostile title cannot break out of its table cell',
           HOSTILE_TITLE in text and text.count('\n| https://site.test/pricing |') == 1)


def test_report_without_a_baseline():
    app = client()
    text = app.get(f'/api/crawls/{FIRST}/report?format=md').data.decode()
    result('report: says so plainly when there is nothing to compare against',
           'No earlier completed crawl' in text)


def test_bad_requests_are_refused():
    app = client()
    status, body = get_json(app, export_url('nonsense'))
    result('unknown dataset: 400 listing the valid ones',
           status == 400 and body.get('success') is False
           and set(body.get('datasets', [])) == set(DATASETS), f'status={status}')

    status, body = get_json(app, export_url('issues', 'parquet'))
    result('unknown format: 400 listing the valid ones',
           status == 400 and body.get('success') is False
           and 'csv' in body.get('formats', []), f'status={status}')

    status, body = get_json(app, f'/api/crawls/{SECOND}/export/all?format=csv')
    result('unknown bundle format: 400, csv is not a bundle',
           status == 400 and body.get('formats') == ['xlsx', 'zip'], f'status={status}')

    status, body = get_json(app, f'/api/crawls/{SECOND}/report?format=pdf')
    result('unknown report format: 400 listing the valid ones',
           status == 400 and body.get('formats') == ['html', 'md'], f'status={status}')


def test_missing_crawl_is_not_found():
    app = client()
    for url in ('/api/crawls/999999/export/issues?format=json',
                '/api/crawls/999999/export/all?format=zip',
                '/api/crawls/999999/report?format=md'):
        status, body = get_json(app, url)
        result(f'missing crawl: 404 from {url.split("/")[-1].split("?")[0]}',
               status == 404 and body.get('success') is False, f'status={status}')


def test_rail_filters_reach_the_export():
    """The filters the UI's rail offers must narrow an export the same way.

    Without these the Export button silently hands back more rows than the
    screen is showing, which is worse than refusing.
    """
    app = client()
    _, ok = get_json(app, export_url('pages', 'json', '&status_class=2xx'))
    _, missing = get_json(app, export_url('pages', 'json', '&status_class=4xx'))
    result('filter status_class: 2xx and 4xx split the page export',
           all(200 <= row['status_code'] < 300 for row in ok['rows'])
           and [row['url'] for row in missing['rows']] == ['https://site.test/gone'],
           f"2xx={ok['count']} 4xx={missing['count']}")

    _, html = get_json(app, export_url('pages', 'json', '&content_type=html'))
    _, images = get_json(app, export_url('pages', 'json', '&content_type=images'))
    result('filter content_type: html and images split the page export',
           images['count'] == 1 and images['rows'][0]['url'].endswith('.png')
           and html['count'] == ok['count'] + missing['count'] - 1,
           f"html={html['count']} images={images['count']}")

    _, crawled_only = get_json(app, export_url('sitemap-coverage', 'json',
                                               '&state=crawled_only'))
    _, both = get_json(app, export_url('sitemap-coverage', 'json', '&state=in_both'))
    result('filter state: the coverage export answers its own column',
           crawled_only['count'] >= 1
           and all(row['state'] == 'crawled_only' for row in crawled_only['rows'])
           and all(row['state'] == 'in_both' for row in both['rows']),
           f"crawled_only={crawled_only['count']} in_both={both['count']}")

    _, unknown = get_json(app, export_url('pages', 'json', '&status_class=9xx'))
    result('filter status_class: an impossible value returns nothing, not everything',
           unknown['count'] == 0, str(unknown['count']))


TESTS = (
    test_rail_filters_reach_the_export,
    test_every_dataset_has_rows,
    test_dataset_columns_are_the_documented_ones,
    test_catalogue_lists_every_dataset,
    test_derived_views_are_built_from_the_stored_rows,
    test_csv_format,
    test_ndjson_format,
    test_field_selection,
    test_xlsx_format,
    test_export_all_workbook,
    test_export_all_zip,
    test_filters_narrow_the_export,
    test_html_report,
    test_markdown_report,
    test_report_without_a_baseline,
    test_bad_requests_are_refused,
    test_missing_crawl_is_not_found,
)


def main_():
    global FIRST, SECOND, ISSUES

    handle, path = tempfile.mkstemp(suffix='.db')
    os.close(handle)
    original = crawl_db.DB_FILE
    try:
        crawl_db.DB_FILE = path
        import main as app_module
        app_module.crawl_db = crawl_db  # the app reads through the same module
        # main.py does not register this blueprint yet, so the test does it
        # itself - before the first request, which is when Flask allows it.
        if 'exports' not in app_module.app.blueprints:
            app_module.app.register_blueprint(create_blueprint(app_module.login_required))
        crawl_db.init_crawl_tables()
        FIRST, SECOND, ISSUES = build_fixture_crawls()

        for test in TESTS:
            print(f'\n{test.__name__}:')
            try:
                test()
            except Exception as exc:  # a raised test is a failed test
                result(test.__name__, False, f'raised {type(exc).__name__}: {exc}')
    finally:
        crawl_db.DB_FILE = original
        os.unlink(path)

    print('\n' + '=' * 64)
    print(f'  {len(PASSED)} passed, {len(FAILED)} failed')
    if FAILED:
        print('  failed: ' + ', '.join(FAILED))
    return 1 if FAILED else 0


if __name__ == '__main__':
    sys.exit(main_())
