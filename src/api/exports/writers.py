"""Serialisation. Every writer yields chunks; none builds one string.

A 500k-row crawl has to leave the process while it is still being produced:
CSV, JSON and NDJSON are yielded row-batch by row-batch, the zip is written
into a spool that is drained between entries, and the workbook goes through
openpyxl's write-only mode (which spills to a temp file) before being streamed
back and deleted.
"""

import csv
import io
import json
import os
import tempfile
import zipfile
from datetime import datetime

from flask import Response, jsonify, request, stream_with_context

from src.api.exports.datasets import DATASETS, filtered

XLSX_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
XLSX_MAX_ROWS = 1_048_576          # Excel's hard limit, header row included
XLSX_CELL_LIMIT = 32767            # and its per-cell character limit
CSV_FLUSH_ROWS = 500               # rows between yields while streaming
FILE_CHUNK = 64 * 1024

ROW_FORMATS = ('csv', 'json', 'ndjson', 'xlsx')
BUNDLE_FORMATS = ('xlsx', 'zip')
REPORT_FORMATS = ('html', 'md')


# ---------------------------------------------------------------------------
# serialisation - every writer yields chunks, never one string
# ---------------------------------------------------------------------------

def _cell(value):
    """Flatten a stored JSON column into something a CSV cell can hold."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str, ensure_ascii=False)
    return value


def _shape(row, fields):
    return {field: _cell(row.get(field)) for field in fields}


def csv_chunks(fields, rows):
    """CSV with a utf-8 BOM, so Excel opens non-ASCII URLs unmangled."""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(fields), extrasaction='ignore',
                            restval='')
    writer.writeheader()
    yield '\ufeff' + _drain(buffer)

    for count, row in enumerate(rows, 1):
        writer.writerow(_shape(row, fields))
        if count % CSV_FLUSH_ROWS == 0:
            yield _drain(buffer)
    tail = _drain(buffer)
    if tail:
        yield tail


def _drain(buffer):
    text = buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)
    return text


def ndjson_chunks(fields, rows):
    for row in rows:
        yield json.dumps(_shape(row, fields), default=str, ensure_ascii=False) + '\n'


def json_chunks(crawl_id, dataset, fields, rows):
    """Streamed object; `count` is written last, once the rows are known."""
    yield ('{"success": true, "crawl_id": %d, "dataset": %s, "rows": ['
           % (crawl_id, json.dumps(dataset)))
    count = 0
    for row in rows:
        yield ('' if count == 0 else ',') + json.dumps(_shape(row, fields),
                                                       default=str, ensure_ascii=False)
        count += 1
    yield '], "count": %d}' % count


def _xlsx_cell(value):
    """Coerce a value into something openpyxl will accept in a cell."""
    if value is None or isinstance(value, bool) or isinstance(value, (int, float)):
        return value
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    text = ILLEGAL_CHARACTERS_RE.sub('', str(_cell(value)))
    return text[:XLSX_CELL_LIMIT]


def _fill_sheet(sheet, fields, rows):
    """Append rows under a header, stopping at Excel's row limit.

    Returns how many rows were dropped. When any were, the last usable row
    says so in the sheet itself - a truncated workbook that looks complete is
    worse than no workbook.
    """
    sheet.append(list(fields))
    room = XLSX_MAX_ROWS - 2          # header, and a line to report truncation
    written = dropped = 0
    for row in rows:
        if written >= room:
            dropped += 1
            continue
        sheet.append([_xlsx_cell(row.get(field)) for field in fields])
        written += 1
    if dropped:
        note = (f'TRUNCATED: {dropped} more rows were left out - a worksheet '
                f'holds at most {XLSX_MAX_ROWS} rows. Export this dataset as '
                f'CSV or NDJSON for all of it.')
        sheet.append([note] + [None] * (len(fields) - 1))
    return dropped


def _save_workbook(workbook):
    """Save to a temp file and hand back its path, for streaming."""
    handle, path = tempfile.mkstemp(suffix='.xlsx', prefix='toadstone-export-')
    os.close(handle)
    try:
        workbook.save(path)
    except Exception:
        os.unlink(path)
        raise
    return path


def workbook_of_one(dataset, fields, rows):
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    _fill_sheet(workbook.create_sheet(title=_sheet_title(dataset)), fields, rows)
    return _save_workbook(workbook)


def workbook_of_everything(crawl_id, crawl):
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    for name, spec in DATASETS.items():
        rows = filtered(spec['loader'](crawl_id, crawl), spec, request.args)
        _fill_sheet(workbook.create_sheet(title=_sheet_title(name)),
                    spec['fields'], rows)
    return _save_workbook(workbook)


def _sheet_title(dataset):
    """Excel sheet names: 31 characters, and none of \\ / ? * [ ] :"""
    title = ''.join(' ' if c in '\\/?*[]:' else c for c in dataset)
    return title[:31] or 'Sheet'


def stream_file(path, mimetype, headers):
    """Stream a temp file back and delete it, however the response ends."""
    def chunks():
        try:
            with open(path, 'rb') as handle:
                while True:
                    chunk = handle.read(FILE_CHUNK)
                    if not chunk:
                        break
                    yield chunk
        finally:
            try:
                os.unlink(path)
            except OSError:
                pass
    return Response(stream_with_context(chunks()), mimetype=mimetype, headers=headers)


class _Spool:
    """A write-only sink zipfile streams into, drained chunk by chunk.

    It has no tell() and no seek(), so zipfile writes data descriptors instead
    of rewinding to patch each entry's header - which is what lets the archive
    leave the process while it is still being built.
    """

    def __init__(self):
        self._parts = []

    def write(self, data):
        self._parts.append(bytes(data))
        return len(data)

    def flush(self):
        pass

    def drain(self):
        chunk = b''.join(self._parts)
        self._parts = []
        return chunk


def zip_of_csvs(crawl_id, crawl):
    """One CSV per dataset, deflated and yielded as the archive is written."""
    spool = _Spool()
    with zipfile.ZipFile(spool, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, spec in DATASETS.items():
            rows = filtered(spec['loader'](crawl_id, crawl), spec, request.args)
            entry_name = _filename('csv', crawl_id, name)
            with archive.open(entry_name, 'w', force_zip64=True) as entry:
                for chunk in csv_chunks(spec['fields'], rows):
                    entry.write(chunk.encode('utf-8'))
                    ready = spool.drain()
                    if ready:
                        yield ready
            ready = spool.drain()
            if ready:
                yield ready
    tail = spool.drain()
    if tail:
        yield tail


def _filename(extension, crawl_id, name):
    return f'crawl-{crawl_id}-{name}-{datetime.now().strftime("%Y-%m-%d")}.{extension}'


def download(extension, crawl_id, name, inline=False):
    disposition = 'inline' if inline else 'attachment'
    return {'Content-Disposition':
            f'{disposition}; filename={_filename(extension, crawl_id, name)}'}


def bad_format(fmt, allowed):
    return jsonify({'success': False, 'error': f'Unknown format {fmt!r}',
                    'formats': list(allowed)}), 400
